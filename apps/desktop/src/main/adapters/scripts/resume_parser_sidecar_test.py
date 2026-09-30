import pathlib
import tempfile
import unittest

from docx import Document
from docx.enum.section import WD_SECTION_START

from resume_parser_sidecar import extract_docx_text, iter_docx_block_text


class DocxBlockTextTests(unittest.TestCase):
    def test_preserves_repeated_role_titles_and_locations_in_distinct_paragraphs(self):
        document = Document()
        lines = [
            "Alex Sample",
            "Senior Product Engineer",
            "Amsterdam Netherlands",
            "EXPERIENCE",
            "Sample Labs",
            "Senior Product Engineer",
            "Amsterdam Netherlands",
            "March 2022 - Present",
        ]
        for line in lines:
            document.add_paragraph(line)

        with tempfile.TemporaryDirectory() as directory:
            source = pathlib.Path(directory) / "synthetic-resume.docx"
            document.save(source)
            text, parser, warnings = extract_docx_text(str(source))

        self.assertEqual(parser, "local_docx")
        self.assertEqual(warnings, [])
        self.assertEqual(text, "\n".join(lines))

    def test_reads_a_linked_header_once_but_preserves_matching_body_text(self):
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Alex Sample"
        document.add_paragraph("Alex Sample")
        document.add_section(WD_SECTION_START.NEW_PAGE)
        self.assertTrue(document.sections[1].header.is_linked_to_previous)

        self.assertEqual(iter_docx_block_text(document).count("Alex Sample"), 2)

    def test_reads_one_merged_cell_once_but_keeps_equal_text_in_distinct_cells(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Senior Engineer"
        table.cell(1, 0).text = "Amsterdam"
        table.cell(1, 1).text = "Amsterdam"

        self.assertEqual(iter_docx_block_text(document), ["Senior Engineer", "Amsterdam", "Amsterdam"])


if __name__ == "__main__":
    unittest.main()
